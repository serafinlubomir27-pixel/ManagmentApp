"""Export projektu — GET /projects/{id}/export/{pdf|csv|xlsx}.

Overujú, že sa vygeneruje platný súbor správneho typu, že sa do neho dostanú
reálne CPM hodnoty a že export rešpektuje prístupové práva k projektu.

Spustenie:
    py -m pytest tests/test_export.py -v
"""
import io
import zipfile

from fastapi.testclient import TestClient

# DB + JWT kľúč pripraví conftest.py.
import backend.main as main
from logic import web_export

client = TestClient(main.app)


def _admin() -> dict:
    r = client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _project_with_tasks(headers: dict) -> int:
    """Projekt A(3d) → B(4d), takže CPM dá A: 0–3, B: 3–7, obe kritické."""
    pid = client.post("/projects/", json={"name": "Export projekt — Žilina"}, headers=headers).json()["id"]
    a = client.post(f"/projects/{pid}/tasks", json={"name": "Príprava", "duration": 3}, headers=headers).json()["id"]
    b = client.post(f"/projects/{pid}/tasks", json={"name": "Montáž", "duration": 4}, headers=headers).json()["id"]
    client.post(f"/tasks/{b}/dependencies", params={"depends_on": a}, headers=headers)
    return pid


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_export_pdf_returns_valid_document():
    h = _admin()
    pid = _project_with_tasks(h)

    r = client.get(f"/projects/{pid}/export/pdf", headers=h)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-"), "odpoveď nie je PDF"
    assert len(r.content) > 1000


def test_export_filename_carries_diacritics():
    """Diakritika prejde cez filename* (RFC 5987), ASCII fallback ju nahradí."""
    h = _admin()
    pid = _project_with_tasks(h)

    r = client.get(f"/projects/{pid}/export/pdf", headers=h)
    disposition = r.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert "%C5%BDilina" in disposition          # "Žilina" percent-encoded
    assert 'filename="Export_projekt_Zilina.pdf"' in disposition
    # Bez expose hlavičky si prehliadač cez CORS názov súboru neprečíta.
    assert "Content-Disposition" in r.headers["access-control-expose-headers"]


# ── CSV ──────────────────────────────────────────────────────────────────────

def test_export_csv_has_bom_and_cpm_values():
    h = _admin()
    pid = _project_with_tasks(h)

    r = client.get(f"/projects/{pid}/export/csv", headers=h)
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"\xef\xbb\xbf"), "chýba UTF-8 BOM — Excel by rozbil diakritiku"

    text = r.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0].startswith("Úloha;Stav;Priorita")

    rows = {line.split(";")[0]: line.split(";") for line in lines[1:]}
    assert set(rows) == {"Príprava", "Montáž"}

    header = lines[0].split(";")
    es, ef, crit = header.index("ES"), header.index("EF"), header.index("Kritická")
    assert rows["Príprava"][es] == "0" and rows["Príprava"][ef] == "3"
    assert rows["Montáž"][es] == "3" and rows["Montáž"][ef] == "7"
    assert rows["Montáž"][crit] == "Áno"


# ── XLSX ─────────────────────────────────────────────────────────────────────

def test_export_xlsx_is_valid_workbook():
    h = _admin()
    pid = _project_with_tasks(h)

    r = client.get(f"/projects/{pid}/export/xlsx", headers=h)
    assert r.status_code == 200, r.text
    # .xlsx je ZIP archív — overíme, že sa dá otvoriť a má obe hárky.
    assert zipfile.is_zipfile(io.BytesIO(r.content))

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Úlohy", "Súhrn"]

    ws = wb["Úlohy"]
    assert ws.max_row == 3                       # hlavička + 2 úlohy
    assert ws.freeze_panes == "A2"
    assert [c.value for c in ws[1]][:3] == ["Úloha", "Stav", "Priorita"]

    summary = {row[0]: row[1] for row in wb["Súhrn"].iter_rows(values_only=True)}
    assert summary["Úloh celkom"] == 2
    assert summary["Trvanie projektu (dni, CPM)"] == 7


# ── Prístupové práva ─────────────────────────────────────────────────────────

def test_export_requires_authentication():
    h = _admin()
    pid = _project_with_tasks(h)
    for fmt in ("pdf", "csv", "xlsx"):
        assert client.get(f"/projects/{pid}/export/{fmt}").status_code == 401


def test_export_rejects_foreign_project():
    """Používateľ z inej organizácie nesmie stiahnuť cudzí projekt."""
    owner = _admin()
    pid = _project_with_tasks(owner)

    signup = client.post("/auth/signup", json={
        "email": "cudzi@example.com",
        "password": "CudzieHeslo123",
        "full_name": "Cudzi Pouzivatel",
        "organization_name": "Cudzia firma",
    })
    assert signup.status_code == 201, signup.text
    outsider = {"Authorization": f"Bearer {signup.json()['access_token']}"}

    r = client.get(f"/projects/{pid}/export/pdf", headers=outsider)
    assert r.status_code == 404, f"cudzí projekt sa dal exportovať ({r.status_code})"


# ── Okrajové prípady generátora ──────────────────────────────────────────────

def test_export_empty_project_does_not_crash():
    h = _admin()
    pid = client.post("/projects/", json={"name": "Prázdny"}, headers=h).json()["id"]

    assert client.get(f"/projects/{pid}/export/pdf", headers=h).content.startswith(b"%PDF-")
    assert client.get(f"/projects/{pid}/export/csv", headers=h).status_code == 200
    assert client.get(f"/projects/{pid}/export/xlsx", headers=h).status_code == 200


def test_gantt_skips_tasks_without_schedule():
    """Úlohy bez CPM rozpisu (ef=0) sa do Ganttu nekreslia."""
    assert web_export._gantt_drawing([{"id": 1, "name": "X", "es": 0, "ef": 0}], 400, "Helvetica", "Helvetica-Bold") is None


def test_gantt_truncates_long_projects():
    """Drawing sa nevie zalomiť na ďalšiu stranu, preto sa počet riadkov orezáva."""
    many = [
        {"id": i, "name": f"Úloha {i}", "es": i, "ef": i + 1,
         "status": "pending", "is_critical": False, "total_float": 0}
        for i in range(web_export.MAX_GANTT_ROWS + 15)
    ]
    d = web_export._gantt_drawing(many, 400, "Helvetica", "Helvetica-Bold")
    assert d is not None
    # Výška zodpovedá orezanému počtu riadkov, nie všetkým úlohám.
    assert d.height < (web_export.MAX_GANTT_ROWS + 15) * 13
