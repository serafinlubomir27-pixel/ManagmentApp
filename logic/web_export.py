"""Export projektu pre webové API — PDF / CSV / XLSX do pamäte (bytes).

Na rozdiel od logic.export_manager (desktop) tento modul nezávisí od matplotlib
a nezapisuje na disk — Gantt sa kreslí priamo primitívami reportlabu a výstup je
BytesIO, ktorý FastAPI streamuje klientovi.
"""
from __future__ import annotations

import csv
import io
import os
from datetime import date

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.graphics.shapes import Drawing, Group, Line, Rect, String
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

_FONT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "fonts")

PRIMARY = colors.HexColor("#1976D2")
DANGER = colors.HexColor("#D32F2F")
SUCCESS = colors.HexColor("#388E3C")
BG_DARK = colors.HexColor("#263238")
BG_ROW = colors.HexColor("#ECEFF1")
BG_CRIT = colors.HexColor("#FFEBEE")
GRID = colors.HexColor("#CFD8DC")
TEXT = colors.HexColor("#212121")
TEXT_DIM = colors.HexColor("#546E7A")

STATUS_SK = {
    "pending": "Čaká",
    "in_progress": "Prebieha",
    "completed": "Hotová",
    "blocked": "Blokovaná",
}
PRIORITY_SK = {
    "low": "Nízka",
    "medium": "Stredná",
    "high": "Vysoká",
    "critical": "Kritická",
}
STATUS_FILL = {
    "pending": colors.HexColor("#B0BEC5"),
    "in_progress": colors.HexColor("#42A5F5"),
    "completed": colors.HexColor("#66BB6A"),
    "blocked": colors.HexColor("#EF5350"),
}

# Gantt sa vykresľuje ako jeden neprerušiteľný Drawing, preto sa dlhé projekty
# orežú — tabuľka úloh nižšie zostáva kompletná.
MAX_GANTT_ROWS = 32


def _register_fonts() -> tuple[str, str]:
    """Zaregistruje DejaVu (plná slovenská diakritika). Fallback: Helvetica."""
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(_FONT_DIR, "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold")
        return "DejaVu", "DejaVu-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def _truthy(value) -> bool:
    """is_critical je 0/1 v SQLite a bool v PostgreSQL."""
    return bool(value)


def _styles(f: str, fb: str) -> dict:
    return {
        "title": ParagraphStyle("title", fontSize=17, fontName=fb, textColor=colors.white, leading=21),
        "subtitle": ParagraphStyle("subtitle", fontSize=9, fontName=f, textColor=colors.HexColor("#B0BEC5")),
        "section": ParagraphStyle("section", fontSize=12, fontName=fb, textColor=PRIMARY, spaceBefore=14, spaceAfter=6),
        "body": ParagraphStyle("body", fontSize=9, fontName=f, textColor=TEXT, leading=13),
        "dim": ParagraphStyle("dim", fontSize=8, fontName=f, textColor=TEXT_DIM, leading=11),
        "th": ParagraphStyle("th", fontSize=7.5, fontName=fb, textColor=colors.white, alignment=TA_CENTER),
        "cell": ParagraphStyle("cell", fontSize=7.5, fontName=f, textColor=TEXT, leading=10),
        "cell_crit": ParagraphStyle("cell_crit", fontSize=7.5, fontName=fb, textColor=DANGER, leading=10),
    }


def _header_banner(project: dict, width: float, st: dict) -> Table:
    name = project.get("name") or "Projekt"
    desc = project.get("description") or ""
    meta = f"Nodus · vygenerované {date.today().strftime('%d.%m.%Y')}"
    inner = [
        [Paragraph(name, st["title"])],
        [Paragraph(desc or meta, st["subtitle"])],
    ]
    if desc:
        inner.append([Paragraph(meta, st["subtitle"])])
    t = Table(inner, colWidths=[width])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BG_DARK),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (0, 0), 12),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 12),
    ]))
    return t


def _stats_row(tasks: list[dict], width: float, f: str, fb: str) -> Table:
    total = len(tasks)
    done = sum(1 for t in tasks if t.get("status") == "completed")
    critical = sum(1 for t in tasks if _truthy(t.get("is_critical")))
    duration = max((t.get("ef") or 0 for t in tasks), default=0)
    progress = round(done / total * 100) if total else 0

    cells = [
        (str(total), "Úloh celkom"),
        (str(done), "Dokončených"),
        (f"{progress}%", "Postup"),
        (str(critical), "Kritických"),
        (f"{duration}d", "Trvanie (CPM)"),
    ]
    col = width / len(cells)
    data = [[v for v, _ in cells], [lbl for _, lbl in cells]]
    t = Table(data, colWidths=[col] * len(cells))
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), fb),
        ("FONTSIZE", (0, 0), (-1, 0), 15),
        ("TEXTCOLOR", (0, 0), (-1, 0), PRIMARY),
        ("FONTNAME", (0, 1), (-1, 1), f),
        ("FONTSIZE", (0, 1), (-1, 1), 7.5),
        ("TEXTCOLOR", (0, 1), (-1, 1), TEXT_DIM),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 10),
        ("BACKGROUND", (0, 0), (-1, -1), BG_ROW),
        ("LINEAFTER", (0, 0), (-2, -1), 0.5, colors.white),
    ]))
    return t


def _gantt_drawing(tasks: list[dict], width: float, f: str, fb: str) -> Drawing | None:
    """Gantt diagram kreslený primitívami reportlabu (bez matplotlib)."""
    scheduled = [t for t in tasks if (t.get("ef") or 0) > 0]
    if not scheduled:
        return None

    scheduled.sort(key=lambda t: ((t.get("es") or 0), t.get("id") or 0))
    truncated = len(scheduled) > MAX_GANTT_ROWS
    rows = scheduled[:MAX_GANTT_ROWS]

    duration = max((t.get("ef") or 0) for t in scheduled) or 1
    label_w = 118.0
    chart_w = width - label_w
    row_h = 13.0
    axis_h = 16.0
    height = len(rows) * row_h + axis_h + 6

    day_w = chart_w / duration
    d = Drawing(width, height)

    # Zvislá mriežka — krok volíme tak, aby nebola prehustená
    step = max(1, duration // 12)
    day = 0
    while day <= duration:
        x = label_w + day * day_w
        d.add(Line(x, axis_h, x, height, strokeColor=GRID, strokeWidth=0.4))
        d.add(String(x, 5, f"d{day}", fontName=f, fontSize=5.5, fillColor=TEXT_DIM, textAnchor="middle"))
        day += step

    for i, t in enumerate(rows):
        y = height - (i + 1) * row_h
        es = t.get("es") or 0
        ef = t.get("ef") or 0
        crit = _truthy(t.get("is_critical"))

        name = t.get("name") or ""
        if len(name) > 26:
            name = name[:25] + "…"
        d.add(String(
            2, y + 3.5, name,
            fontName=fb if crit else f, fontSize=6,
            fillColor=DANGER if crit else TEXT,
        ))

        bar_x = label_w + es * day_w
        bar_w = max(1.5, (ef - es) * day_w)
        fill = DANGER if crit else STATUS_FILL.get(t.get("status"), colors.HexColor("#42A5F5"))
        d.add(Rect(
            bar_x, y + 1.5, bar_w, row_h - 4.5,
            fillColor=fill, strokeColor=None, rx=1.5, ry=1.5,
        ))

        # Rezerva (total float) ako svetlá stopa za úlohou
        float_days = t.get("total_float") or 0
        if float_days > 0 and not crit:
            d.add(Rect(
                bar_x + bar_w, y + 3.5, float_days * day_w, row_h - 8.5,
                fillColor=colors.HexColor("#E0E0E0"), strokeColor=None,
            ))

    if truncated:
        d.add(String(
            2, 5, f"(zobrazených prvých {MAX_GANTT_ROWS} z {len(scheduled)} úloh)",
            fontName=f, fontSize=5.5, fillColor=TEXT_DIM,
        ))

    return d


def _gantt_legend(f: str) -> Drawing:
    items = [
        ("Kritická cesta", DANGER),
        ("Prebieha", STATUS_FILL["in_progress"]),
        ("Hotová", STATUS_FILL["completed"]),
        ("Čaká", STATUS_FILL["pending"]),
        ("Rezerva", colors.HexColor("#E0E0E0")),
    ]
    d = Drawing(460, 12)
    x = 0.0
    for label, col in items:
        d.add(Rect(x, 2, 7, 7, fillColor=col, strokeColor=None, rx=1, ry=1))
        d.add(String(x + 10, 3.5, label, fontName=f, fontSize=6.5, fillColor=TEXT_DIM))
        x += 12 + len(label) * 3.9 + 12
    return d


def _task_table(tasks: list[dict], logged: dict[int, float], width: float, st: dict) -> Table:
    headers = ["Úloha", "Stav", "Priorita", "Trv.", "ES", "EF", "LS", "LF", "Rez.", "Hodiny"]
    data = [[Paragraph(h, st["th"]) for h in headers]]

    for t in tasks:
        crit = _truthy(t.get("is_critical"))
        style = st["cell_crit"] if crit else st["cell"]
        name = t.get("name") or ""
        hours = logged.get(t.get("id"), 0.0)
        est = t.get("estimated_hours")
        hours_txt = f"{hours:g}" + (f"/{est:g}" if est else "")
        data.append([
            Paragraph(("● " if crit else "") + name, style),
            Paragraph(STATUS_SK.get(t.get("status"), t.get("status") or "—"), style),
            Paragraph(PRIORITY_SK.get(t.get("priority"), t.get("priority") or "—"), style),
            Paragraph(f"{t.get('duration') or 0}d", style),
            Paragraph(str(t.get("es") or 0), style),
            Paragraph(str(t.get("ef") or 0), style),
            Paragraph(str(t.get("ls") or 0), style),
            Paragraph(str(t.get("lf") or 0), style),
            Paragraph(f"{t.get('total_float') or 0}d", style),
            Paragraph(hours_txt, style),
        ])

    # Stav/Priorita potrebujú viac miesta než číselné CPM stĺpce, inak sa slová lámu.
    fixed = [50, 48, 30, 24, 24, 24, 24, 32, 42]
    col_widths = [width - sum(fixed)] + fixed
    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), BG_DARK),
        ("GRID", (0, 0), (-1, -1), 0.3, GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (3, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    for i, task in enumerate(tasks, start=1):
        if _truthy(task.get("is_critical")):
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), BG_CRIT))
        elif i % 2 == 0:
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")))
    t.setStyle(TableStyle(style_cmds))
    return t


def _footer(canvas, doc, f: str, project_name: str) -> None:
    canvas.saveState()
    canvas.setFont(f, 7)
    canvas.setFillColor(TEXT_DIM)
    canvas.drawString(1.8 * cm, 1.1 * cm, f"Nodus · {project_name}")
    canvas.drawRightString(A4[0] - 1.8 * cm, 1.1 * cm, f"Strana {doc.page}")
    canvas.setStrokeColor(GRID)
    canvas.setLineWidth(0.4)
    canvas.line(1.8 * cm, 1.4 * cm, A4[0] - 1.8 * cm, 1.4 * cm)
    canvas.restoreState()


def build_project_pdf(project: dict, tasks: list[dict], time_summary: list[dict] | None = None) -> bytes:
    """Vygeneruje PDF report projektu a vráti ho ako bajty."""
    f, fb = _register_fonts()
    st = _styles(f, fb)
    width = A4[0] - 3.6 * cm
    logged = {r["id"]: float(r.get("logged_hours") or 0) for r in (time_summary or [])}

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        topMargin=1.5 * cm, bottomMargin=1.8 * cm,
        title=f"Nodus — {project.get('name', 'Projekt')}",
        author="Nodus",
    )

    story: list = [
        _header_banner(project, width, st),
        Spacer(1, 10),
        _stats_row(tasks, width, f, fb),
    ]

    gantt = _gantt_drawing(tasks, width, f, fb)
    if gantt is not None:
        story += [
            Paragraph("Ganttov diagram", st["section"]),
            gantt,
            Spacer(1, 4),
            _gantt_legend(f),
        ]

    critical = [t for t in tasks if _truthy(t.get("is_critical"))]
    if critical:
        chain = "  →  ".join(t.get("name") or "" for t in sorted(critical, key=lambda x: x.get("es") or 0))
        story += [
            Paragraph("Kritická cesta", st["section"]),
            Paragraph(chain, st["body"]),
            Spacer(1, 2),
            Paragraph(
                f"{len(critical)} úloh bez časovej rezervy — akékoľvek ich omeškanie posunie termín projektu.",
                st["dim"],
            ),
        ]

    if tasks:
        story += [
            Paragraph("Zoznam úloh", st["section"]),
            _task_table(tasks, logged, width, st),
        ]
    else:
        story += [Paragraph("Projekt zatiaľ neobsahuje žiadne úlohy.", st["body"])]

    story += [
        Spacer(1, 14),
        HRFlowable(width="100%", thickness=0.4, color=GRID),
        Spacer(1, 4),
        Paragraph(
            "ES = najskorší začiatok, EF = najskorší koniec, LS = najneskorší začiatok, "
            "LF = najneskorší koniec, Rezerva = celková časová rezerva (total float). "
            "Hodnoty vychádzajú z metódy kritickej cesty (CPM).",
            st["dim"],
        ),
    ]

    def _page(canvas, d):
        _footer(canvas, d, f, project.get("name") or "Projekt")

    doc.build(story, onFirstPage=_page, onLaterPages=_page)
    return buf.getvalue()


CSV_COLUMNS = [
    ("name", "Úloha"),
    ("status", "Stav"),
    ("priority", "Priorita"),
    ("assigned_username", "Pridelený"),
    ("due_date", "Termín"),
    ("duration", "Trvanie (dni)"),
    ("delay_days", "Omeškanie (dni)"),
    ("es", "ES"),
    ("ef", "EF"),
    ("ls", "LS"),
    ("lf", "LF"),
    ("total_float", "Rezerva"),
    ("is_critical", "Kritická"),
    ("estimated_hours", "Odhad (h)"),
    ("logged_hours", "Odpracované (h)"),
    ("category", "Kategória"),
    ("description", "Popis"),
]


def _export_rows(tasks: list[dict], time_summary: list[dict] | None) -> list[list]:
    logged = {r["id"]: float(r.get("logged_hours") or 0) for r in (time_summary or [])}
    rows = []
    for t in tasks:
        row = []
        for key, _ in CSV_COLUMNS:
            if key == "logged_hours":
                row.append(logged.get(t.get("id"), 0))
            elif key == "is_critical":
                row.append("Áno" if _truthy(t.get("is_critical")) else "Nie")
            elif key == "status":
                row.append(STATUS_SK.get(t.get("status"), t.get("status") or ""))
            elif key == "priority":
                row.append(PRIORITY_SK.get(t.get("priority"), t.get("priority") or ""))
            else:
                row.append(t.get(key) if t.get(key) is not None else "")
        rows.append(row)
    return rows


def build_tasks_csv(tasks: list[dict], time_summary: list[dict] | None = None) -> bytes:
    """CSV so slovenskými hlavičkami. UTF-8 BOM, aby sa správne otvoril v Exceli."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([label for _, label in CSV_COLUMNS])
    writer.writerows(_export_rows(tasks, time_summary))
    return buf.getvalue().encode("utf-8-sig")


def build_tasks_xlsx(project: dict, tasks: list[dict], time_summary: list[dict] | None = None) -> bytes:
    """Excel zošit s formátovanou hlavičkou a zvýraznenou kritickou cestou."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Úlohy"

    header_fill = PatternFill("solid", fgColor="263238")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    crit_fill = PatternFill("solid", fgColor="FFEBEE")

    labels = [label for _, label in CSV_COLUMNS]
    ws.append(labels)
    for col in range(1, len(labels) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_vals, task in zip(_export_rows(tasks, time_summary), tasks):
        ws.append(row_vals)
        if _truthy(task.get("is_critical")):
            for col in range(1, len(labels) + 1):
                ws.cell(row=ws.max_row, column=col).fill = crit_fill

    widths = {"Úloha": 34, "Popis": 40, "Pridelený": 16, "Kategória": 16}
    for idx, label in enumerate(labels, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(label, 12)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{ws.max_row}"

    # Druhý hárok — súhrn projektu
    meta = wb.create_sheet("Súhrn")
    done = sum(1 for t in tasks if t.get("status") == "completed")
    duration = max((t.get("ef") or 0 for t in tasks), default=0)
    summary_rows = [
        ("Projekt", project.get("name") or ""),
        ("Popis", project.get("description") or ""),
        ("Exportované", date.today().strftime("%d.%m.%Y")),
        ("Úloh celkom", len(tasks)),
        ("Dokončených", done),
        ("Postup (%)", round(done / len(tasks) * 100) if tasks else 0),
        ("Kritických úloh", sum(1 for t in tasks if _truthy(t.get("is_critical")))),
        ("Trvanie projektu (dni, CPM)", duration),
    ]
    for label, value in summary_rows:
        meta.append([label, value])
    for row in range(1, len(summary_rows) + 1):
        meta.cell(row=row, column=1).font = Font(bold=True)
    meta.column_dimensions["A"].width = 30
    meta.column_dimensions["B"].width = 42

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
